import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUTPUT_PATH = r"D:\new web record file\Smart_Wireless_Ultrasonic_Thickness_Gauging_Project_Report_Rohini_College.docx"
PDF_OUTPUT_PATH = r"D:\new web record file\Smart_Wireless_Ultrasonic_Thickness_Gauging_Project_Report_Rohini_College.pdf"

DOCX_COPY_PATH = r"C:\final doc\Smart_Wireless_Ultrasonic_Thickness_Gauging_Project_Report_Rohini_College.docx"
PDF_COPY_PATH = r"C:\final doc\Smart_Wireless_Ultrasonic_Thickness_Gauging_Project_Report_Rohini_College.pdf"

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_full_grid_borders(table, color="B0BEC5", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_no_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header_footer(doc):
    for s_idx, section in enumerate(doc.sections):
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True

        # Header for pages 2+
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.paragraph_format.space_after = Pt(2)
        r_hdr = p_hdr.add_run("ROHINI COLLEGE OF ENGINEERING & TECHNOLOGY  |  DEPARTMENT OF MECHANICAL ENGINEERING")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(100, 116, 139)

        # Footer for pages 2+
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # We can add page field XML or handle dynamically

def build_report():
    doc = Document()
    add_header_footer(doc)

    PRIMARY_COLOR = RGBColor(15, 23, 42)     # Slate 900
    NAVY_BLUE = RGBColor(15, 76, 129)        # Royal Navy Blue (matches Bonafide screenshot)
    DARK_TEXT = RGBColor(30, 41, 59)         # Slate 800

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = DARK_TEXT
    style_normal.paragraph_format.line_spacing = 1.35
    style_normal.paragraph_format.space_after = Pt(4)

    logo_path = os.path.abspath("rcet_college_logo_0.png")

    def page_break():
        doc.add_page_break()

    def add_page_title(title, space_before=16, space_after=16, color=PRIMARY_COLOR, font_size=15):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.name = "Times New Roman"
        r.font.color.rgb = color
        return p

    def add_chapter_head(num, title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run(f"CHAPTER {num} — {title.upper()}")
        r.bold = True
        r.font.size = Pt(14)
        r.font.name = "Times New Roman"
        r.font.color.rgb = PRIMARY_COLOR
        return p

    def add_sec_head(title, space_before=10, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(12.5)
        r.font.name = "Times New Roman"
        r.font.color.rgb = PRIMARY_COLOR
        return p

    def add_p(text, space_after=4, line_spacing=1.35, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11.5)
        return p

    def add_bullet(text, space_after=2):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.30
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11.5)
        return p

    # =========================================================================
    # PAGE 1: FRONT COVER PAGE (College Official Thesis Format)
    # =========================================================================
    p_c1 = doc.add_paragraph()
    p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c1.paragraph_format.space_before = Pt(0)
    p_c1.paragraph_format.space_after = Pt(1)
    r = p_c1.add_run("ROHINI COLLEGE OF ENGINEERING AND TECHNOLOGY")
    r.bold = True
    r.font.size = Pt(13)
    r.font.name = "Times New Roman"
    r.font.color.rgb = PRIMARY_COLOR

    p_c2 = doc.add_paragraph()
    p_c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c2.paragraph_format.space_before = Pt(0)
    p_c2.paragraph_format.space_after = Pt(4)
    r = p_c2.add_run("Palkulam, Kanyakumari – 629401\nDEPARTMENT OF MECHANICAL ENGINEERING")
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = "Times New Roman"

    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(2)
        p_logo.paragraph_format.space_after = Pt(6)
        p_logo.add_run().add_picture(logo_path, width=Inches(1.4))

    p_c3 = doc.add_paragraph()
    p_c3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c3.paragraph_format.space_before = Pt(2)
    p_c3.paragraph_format.space_after = Pt(4)
    r = p_c3.add_run("A PROJECT REPORT ON")
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = "Times New Roman"

    p_c4 = doc.add_paragraph()
    p_c4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c4.paragraph_format.space_before = Pt(2)
    p_c4.paragraph_format.space_after = Pt(8)
    r = p_c4.add_run("DESIGN AND DEVELOPMENT OF A SMART WIRELESS\nULTRASONIC THICKNESS GAUGING SYSTEM WITH A WIRELESS PROBE")
    r.bold = True
    r.font.size = Pt(13)
    r.font.name = "Times New Roman"
    r.font.color.rgb = NAVY_BLUE

    p_c5 = doc.add_paragraph()
    p_c5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c5.paragraph_format.space_before = Pt(0)
    p_c5.paragraph_format.space_after = Pt(2)
    r = p_c5.add_run("Submitted by")
    r.font.size = Pt(10.5)
    r.font.name = "Times New Roman"

    students = [
        ("BHUVANESH RS", "Register No. 963323114017"),
        ("MESHOCK M", "Register No. 963323114037"),
        ("RICKSON S", "Register No. 963323114049"),
        ("SAJIN RAJ T", "Register No. 963323114051")
    ]
    for sname, sreg in students:
        p_s = doc.add_paragraph()
        p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_s.paragraph_format.space_before = Pt(0)
        p_s.paragraph_format.space_after = Pt(0.5)
        r = p_s.add_run(f"{sname}   ({sreg})")
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.name = "Times New Roman"

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_before = Pt(4)
    p_deg.paragraph_format.space_after = Pt(4)
    r = p_deg.add_run("in partial fulfilment for the award of the degree of\nBACHELOR OF ENGINEERING\nIN\nMECHANICAL ENGINEERING")
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = "Times New Roman"

    p_c6 = doc.add_paragraph()
    p_c6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c6.paragraph_format.space_before = Pt(2)
    p_c6.paragraph_format.space_after = Pt(1)
    r = p_c6.add_run("Under the guidance of\n")
    r.font.size = Pt(10)
    r.font.name = "Times New Roman"
    r2 = p_c6.add_run("Dr. S. KAILAINATHAN, M.E., Ph.D.\nProfessor & Head of the Department, Department of Mechanical Engineering")
    r2.bold = True
    r2.font.size = Pt(10.5)
    r2.font.name = "Times New Roman"

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(6)
    p_foot.paragraph_format.space_after = Pt(0)
    r = p_foot.add_run("ROHINI COLLEGE OF ENGINEERING & TECHNOLOGY, PALKULAM\n(AUTONOMOUS)\nANNA UNIVERSITY :: CHENNAI 600 025\nACADEMIC YEAR 2026–2027")
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = "Times New Roman"

    page_break()

    # =========================================================================
    # PAGE 2: BONAFIDE CERTIFICATE (Exact Screenshot Format)
    # =========================================================================
    add_page_title("BONAFIDE CERTIFICATE", space_before=16, space_after=16, color=NAVY_BLUE, font_size=14.5)

    p_b1 = doc.add_paragraph()
    p_b1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_b1.paragraph_format.line_spacing = 1.35
    p_b1.paragraph_format.space_after = Pt(12)
    r = p_b1.add_run(
        "Certified that this project report entitled \"Design and Development of a Smart Wireless "
        "Ultrasonic Thickness Gauging System with a Wireless Probe\" is the bonafide work of "
        "BHUVANESH RS (963323114017), MESHOCK M (963323114037), RICKSON S (963323114049) and "
        "SAJIN RAJ T (963323114051) of the Department of Mechanical Engineering, Rohini College of "
        "Engineering and Technology, who carried out the project work under my supervision."
    )
    r.font.size = Pt(11.5)
    r.font.name = "Times New Roman"

    p_b2 = doc.add_paragraph()
    p_b2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_b2.paragraph_format.line_spacing = 1.35
    p_b2.paragraph_format.space_after = Pt(28)
    r = p_b2.add_run(
        "Certified further that, to the best of my knowledge, the work reported herein does not form "
        "part of any other project report or dissertation on the basis of which a degree or award was "
        "conferred on an earlier occasion on this or any other candidate."
    )
    r.font.size = Pt(11.5)
    r.font.name = "Times New Roman"

    # Guide & HOD 2-column signature layout
    tbl_sig = doc.add_table(rows=1, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_no_borders(tbl_sig)
    tbl_sig.rows[0].cells[0].width = Inches(3.4)
    tbl_sig.rows[0].cells[1].width = Inches(3.4)

    # Left: Project Guide
    c_left = tbl_sig.rows[0].cells[0]
    p_gl = c_left.paragraphs[0]
    p_gl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_gl.paragraph_format.line_spacing = 1.15
    r = p_gl.add_run("Dr. S. Kailainathan\n")
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.name = "Times New Roman"
    r2 = p_gl.add_run("Project Guide\nProfessor, Department of Mechanical Engineering")
    r2.font.size = Pt(10.5)
    r2.font.name = "Times New Roman"

    # Right: HOD
    c_right = tbl_sig.rows[0].cells[1]
    p_gr = c_right.paragraphs[0]
    p_gr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_gr.paragraph_format.line_spacing = 1.15
    r = p_gr.add_run("Head of the Department\n")
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.name = "Times New Roman"
    r2 = p_gr.add_run("Department of Mechanical Engineering\nRohini College of Engineering & Technology")
    r2.font.size = Pt(10.5)
    r2.font.name = "Times New Roman"

    p_viva = doc.add_paragraph()
    p_viva.paragraph_format.space_before = Pt(32)
    p_viva.paragraph_format.space_after = Pt(32)
    r = p_viva.add_run("Submitted for the project viva-voce examination held on ____________________")
    r.font.size = Pt(11)
    r.font.name = "Times New Roman"

    # Internal & External Examiners 2-column signature layout
    tbl_exam = doc.add_table(rows=1, cols=2)
    tbl_exam.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_no_borders(tbl_exam)
    tbl_exam.rows[0].cells[0].width = Inches(3.4)
    tbl_exam.rows[0].cells[1].width = Inches(3.4)

    p_el = tbl_exam.rows[0].cells[0].paragraphs[0]
    p_el.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_el.add_run("Internal Examiner")
    r.bold = True
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"

    p_er = tbl_exam.rows[0].cells[1].paragraphs[0]
    p_er.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_er.add_run("External Examiner")
    r.bold = True
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"

    page_break()

    # =========================================================================
    # PAGE 3: ACKNOWLEDGEMENT
    # =========================================================================
    add_page_title("ACKNOWLEDGEMENT", space_before=18, space_after=18)

    add_p("We acknowledge with gratitude all those who helped us to complete this project successfully. First, we thank the Almighty God for giving us the strength, health and ability to carry out this work.")
    add_p("We express our sincere thanks to the Chairman, Shri K. Neela Marthandan, for providing the facilities and support needed for academic and project activities.")
    add_p("We thank the Pro-Chairman, Dr. N. Neelavishnu, and the Managing Director, Dr. V. M. Blessy Geo, for their encouragement and support to the institution.")
    add_p("We express our gratitude to our Principal, Dr. R. Rajesh, for his encouragement and for supporting the academic development of students.")
    add_p("We extend our sincere thanks to our Head of the Department and Project Guide, Dr. S. Kailainathan, for his valuable guidance, technical suggestions, motivation and supervision throughout this project.")
    add_p("We also thank the faculty members and laboratory staff of the Department of Mechanical Engineering for their assistance during the planning, assembly and documentation stages.")
    add_p("Finally, we are grateful to our parents and friends for their constant encouragement and support, which helped us complete the project work.")

    page_break()

    # =========================================================================
    # PAGE 4: ABSTRACT
    # =========================================================================
    add_page_title("ABSTRACT", space_before=18, space_after=18)

    add_p("This project presents the design and development of a wireless prototype for estimating object thickness using ultrasonic distance sensing and embedded communication. An HC-SR04 sensor is mounted at a fixed reference height of 14.1 cm. A NodeMCU ESP8266 sender measures the air gap between the sensor and the upper surface of an object and calculates an estimated thickness using the relationship: estimated thickness = reference height − measured distance.")
    add_p("The calculated value is transmitted to a receiver NodeMCU using ESP-NOW. A 16×2 I2C LCD provides local display, while the Blynk platform is used for remote monitoring. The project setup also includes power regulation, an indicator LED, and a separate sensing arrangement.")
    add_p("The prototype demonstrates integration of sensing, calculation by the NodeMCU, wireless communication and local/remote display. An important technical limitation is that the HC-SR04 measures distance through air; it is not a conventional ultrasonic material-thickness probe. Therefore, the present implementation is best described as a distance-based thickness-estimation prototype. Accuracy and consistency of repeated readings must be established by testing against known known-thickness samples before numerical performance claims are made.")
    add_p("Keywords: ultrasonic distance sensing, thickness estimation, NodeMCU ESP8266, ESP-NOW, LCD, Blynk IoT, wireless monitoring.")

    page_break()

    # =========================================================================
    # PAGE 5: TABLE OF CONTENTS
    # =========================================================================
    add_page_title("TABLE OF CONTENTS", space_before=18, space_after=14)

    toc_items = [
        ("Front matter", "Bonafide Certificate, Acknowledgement, Abstract (pp. 2–4)"),
        ("Chapter 1", "Introduction, problem statement, aim and objectives (pp. 7–12)"),
        ("Chapter 2", "Literature review and base paper (pp. 13–19)"),
        ("Chapter 3", "System requirements and design (pp. 20–25)"),
        ("Chapter 4", "Hardware and software description (pp. 26–30)"),
        ("Chapter 5", "Methodology and implementation (pp. 31–34)"),
        ("Chapter 6", "Experimental setup and testing (pp. 35–38)"),
        ("Chapter 7", "Results and discussion (pp. 39–41)"),
        ("Chapter 8", "Conclusion and future scope (p. 42)"),
        ("—", "References (p. 43)"),
        ("Appendix A", "Prototype photographs (pp. 43–44)"),
        ("Appendix B", "Final submission checklist (p. 45)")
    ]

    tbl_toc = doc.add_table(rows=len(toc_items)+1, cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_full_grid_borders(tbl_toc)
    tbl_toc.rows[0].cells[0].width = Inches(2.2)
    tbl_toc.rows[0].cells[1].width = Inches(4.5)

    tbl_toc.rows[0].cells[0].paragraphs[0].text = "Section"
    tbl_toc.rows[0].cells[0].paragraphs[0].runs[0].bold = True
    tbl_toc.rows[0].cells[1].paragraphs[0].text = "Title"
    tbl_toc.rows[0].cells[1].paragraphs[0].runs[0].bold = True
    set_cell_background(tbl_toc.rows[0].cells[0], "E2E8F0")
    set_cell_background(tbl_toc.rows[0].cells[1], "E2E8F0")

    for idx, (sec, tit) in enumerate(toc_items):
        r_c = tbl_toc.rows[idx+1]
        r_c.cells[0].paragraphs[0].text = sec
        r_c.cells[0].paragraphs[0].runs[0].bold = True if "Chapter" in sec or "Appendix" in sec else False
        r_c.cells[1].paragraphs[0].text = tit
        set_cell_margins(r_c.cells[0], top=80, bottom=80, left=120, right=120)
        set_cell_margins(r_c.cells[1], top=80, bottom=80, left=120, right=120)

    p_ref_note = doc.add_paragraph()
    p_ref_note.paragraph_format.space_before = Pt(14)
    r = p_ref_note.add_run("Page references above correspond to the current 45-page PDF layout.")
    r.italic = True
    r.font.size = Pt(10.5)

    page_break()

    # =========================================================================
    # PAGE 6: LIST OF ABBREVIATIONS
    # =========================================================================
    add_page_title("LIST OF ABBREVIATIONS", space_before=18, space_after=14)

    abbrs = [
        ("HC-SR04", "Ultrasonic distance measurement module"),
        ("ESP8266", "Wi-Fi-enabled microcontroller system-on-chip"),
        ("ESP-NOW", "Espressif direct wireless communication protocol"),
        ("I2C", "Inter-Integrated Circuit serial communication bus"),
        ("LCD", "Liquid Crystal Display"),
        ("IoT", "Internet of Things"),
        ("MCU", "Microcontroller Unit"),
        ("NDT", "Non-Destructive Testing"),
        ("UTG", "Ultrasonic Thickness Gauging"),
        ("GPIO", "General-Purpose Input/Output"),
        ("IDE", "Integrated Development Environment"),
        ("H", "Reference height above the reference plane"),
        ("d", "Measured distance from sensor to object surface"),
        ("t", "Estimated object thickness")
    ]

    tbl_ab = doc.add_table(rows=len(abbrs)+1, cols=2)
    tbl_ab.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_full_grid_borders(tbl_ab)
    tbl_ab.rows[0].cells[0].width = Inches(2.2)
    tbl_ab.rows[0].cells[1].width = Inches(4.5)

    tbl_ab.rows[0].cells[0].paragraphs[0].text = "Abbreviation"
    tbl_ab.rows[0].cells[0].paragraphs[0].runs[0].bold = True
    tbl_ab.rows[0].cells[1].paragraphs[0].text = "Meaning"
    tbl_ab.rows[0].cells[1].paragraphs[0].runs[0].bold = True
    set_cell_background(tbl_ab.rows[0].cells[0], "E2E8F0")
    set_cell_background(tbl_ab.rows[0].cells[1], "E2E8F0")

    for idx, (ab, mean) in enumerate(abbrs):
        r_c = tbl_ab.rows[idx+1]
        r_c.cells[0].paragraphs[0].text = ab
        r_c.cells[0].paragraphs[0].runs[0].bold = True
        r_c.cells[1].paragraphs[0].text = mean
        set_cell_margins(r_c.cells[0], top=70, bottom=70, left=120, right=120)
        set_cell_margins(r_c.cells[1], top=70, bottom=70, left=120, right=120)

    page_break()

    # =========================================================================
    # PAGE 7: CHAPTER 1 — INTRODUCTION
    # =========================================================================
    add_chapter_head("1", "Introduction")
    add_sec_head("1.1 Background")
    add_p("Thickness measurement is important in manufacturing, maintenance, quality control and product development. The thickness of a component influences its dimensions, fit, mass, structural behaviour and service life. The appropriate measurement method depends on the material, geometry, accessibility, surface condition and required accuracy.")
    add_p("Common contact instruments include vernier calipers and micrometers. For some components, direct access to both sides may be inconvenient. Electronic sensing systems can collect a measurement and make the resulting value available through a local display or a remote interface.")
    add_sec_head("1.2 Project context")
    add_p("This project integrates an HC-SR04 ultrasonic distance module with NodeMCU ESP8266 controllers. The sender calculates an estimate from a fixed reference height, and the receiver presents the value on an LCD. A Blynk interface is included for remote monitoring.")
    page_break()

    # =========================================================================
    # PAGE 8: 1.3 PROBLEM STATEMENT
    # =========================================================================
    add_sec_head("1.3 PROBLEM STATEMENT", space_before=16, space_after=8)
    add_p("A setup with only a local display requires the operator to remain close to the display to read the value. Manual transcription can also introduce mistakes. A wireless link between the sensing and display units provides a way to view the reading from a separate location.")
    add_sec_head("Problem addressed")
    add_bullet("Measure the air distance between the sensor and the upper surface of an object.")
    add_bullet("Estimate thickness from a known reference height.")
    add_bullet("Transmit the calculated value to a separate receiver unit.")
    add_bullet("Display the value locally and through a remote monitoring interface.")
    add_bullet("Document the measurement limitations and the need for checking readings against known values.")
    add_sec_head("Technical boundary")
    add_p("The HC-SR04 measures the distance to an external surface through air. It does not measure the time taken by an ultrasonic wave to travel through the solid material. This distinction must be stated clearly when explaining the current prototype.")
    page_break()

    # =========================================================================
    # PAGE 9: 1.4 AIM AND OBJECTIVES
    # =========================================================================
    add_sec_head("1.4 AIM AND OBJECTIVES", space_before=16, space_after=8)
    add_sec_head("Aim")
    add_p("To develop a wireless prototype for thickness estimation using measured distance with local display and IoT-based remote monitoring.")
    add_sec_head("Objectives")
    add_bullet("1. Interface the HC-SR04 distance sensor with a NodeMCU ESP8266 sender.")
    add_bullet("2. Calculate estimated thickness using a fixed reference height of 14.1 cm.")
    add_bullet("3. Transmit the calculated value to a receiver NodeMCU through ESP-NOW.")
    add_bullet("4. Display the received value on a 16×2 I2C LCD.")
    add_bullet("5. Enable remote viewing of measurement data through Blynk.")
    add_bullet("6. Prepare a test procedure to compare estimates with known known-thickness samples.")
    page_break()

    # =========================================================================
    # PAGE 10: 1.5 SCOPE OF THE PROJECT
    # =========================================================================
    add_sec_head("1.5 SCOPE OF THE PROJECT", space_before=16, space_after=8)
    add_sec_head("Included scope")
    add_bullet("Prototype-level distance measurement and thickness estimation using the reference height.")
    add_bullet("Wireless data transfer between sender and receiver units.")
    add_bullet("Local display of calculated values.")
    add_bullet("Remote viewing through Blynk when connectivity is available.")
    add_bullet("Functional checks and comparison with known-thickness samples.")
    add_sec_head("Excluded from the present scope")
    add_bullet("Certified industrial wall-thickness inspection.")
    add_bullet("Corrosion assessment or acceptance decisions for pressure equipment.")
    add_bullet("High-temperature or hazardous-area operation.")
    add_bullet("Claims of ultrasonic wave propagation through the solid material using the HC-SR04.")
    page_break()

    # =========================================================================
    # PAGE 11: 1.6 SIGNIFICANCE AND REPORT ORGANISATION
    # =========================================================================
    add_sec_head("1.6 SIGNIFICANCE AND REPORT ORGANISATION", space_before=16, space_after=8)
    add_sec_head("Significance")
    add_p("The project provides practical experience in sensor interfacing, microcontroller programming, wireless communication, display integration and IoT monitoring. It also introduces engineering practices such as reference measurement, checking readings against known values, error calculation and documentation of limitations.")
    add_sec_head("Report organisation")
    p_org = doc.add_paragraph()
    p_org.paragraph_format.line_spacing = 1.45
    p_org.paragraph_format.space_before = Pt(6)
    r = p_org.add_run(
        "Chapter 1 introduces the project. Chapter 2 reviews related research. Chapter 3 describes "
        "the system requirements and architecture. Chapter 4 explains the hardware and software. "
        "Chapter 5 describes implementation. Chapter 6 presents the experimental procedure. Chapter 7 "
        "discusses the functional outcome and limitations. Chapter 8 concludes the work and identifies "
        "future improvements."
    )
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    page_break()

    # =========================================================================
    # PAGE 12: 1.7 ENGINEERING REQUIREMENTS
    # =========================================================================
    add_sec_head("1.7 ENGINEERING REQUIREMENTS", space_before=16, space_after=8)
    eng_reqs = [
        ("Sensing", "Read distance to an object surface", "Observe serial output and repeat readings"),
        ("Calculation", "Subtract distance from reference height", "Check calculation against manual arithmetic"),
        ("Wireless link", "Send a measurement packet", "Confirm receiver receives the value"),
        ("Local display", "Show the received estimate", "Compare LCD with sender output"),
        ("Remote monitoring", "Show reading in Blynk", "Confirm dashboard update"),
        ("Power", "Supply suitable regulated voltage", "Measure supply with a multimeter")
    ]
    tbl_req = doc.add_table(rows=len(eng_reqs)+1, cols=3)
    tbl_req.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_full_grid_borders(tbl_req)
    tbl_req.rows[0].cells[0].width = Inches(1.8)
    tbl_req.rows[0].cells[1].width = Inches(2.6)
    tbl_req.rows[0].cells[2].width = Inches(2.3)

    for c_i, h_txt in enumerate(["Requirement", "Description", "Verification method"]):
        cell = tbl_req.rows[0].cells[c_i]
        cell.paragraphs[0].text = h_txt
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "E2E8F0")

    for idx, (rq, desc, vm) in enumerate(eng_reqs):
        r_c = tbl_req.rows[idx+1]
        r_c.cells[0].paragraphs[0].text = rq
        r_c.cells[0].paragraphs[0].runs[0].bold = True
        r_c.cells[1].paragraphs[0].text = desc
        r_c.cells[2].paragraphs[0].text = vm
        for cell in r_c.cells:
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)

    p_eng_sub = doc.add_paragraph()
    p_eng_sub.paragraph_format.space_before = Pt(12)
    p_eng_sub.paragraph_format.line_spacing = 1.35
    r = p_eng_sub.add_run("The system should be assessed one subsystem at a time before full integration. This helps isolate sensor, firmware, communication, display and power issues.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    page_break()

    # =========================================================================
    # PAGE 13: CHAPTER 2 — LITERATURE REVIEW
    # =========================================================================
    add_chapter_head("2", "Literature Review")
    add_p("A literature review identifies established methods, system layouts and technical constraints relevant to a project. For this work, the main themes are ultrasonic thickness measurement, signal processing, wireless wall-thickness monitoring and remote display of measurement data.")
    add_p("The papers discussed in this chapter use specialised ultrasonic measurement methods. They provide research context, but their sensor arrangements are not identical to the HC-SR04 distance-based prototype. The difference is important when defining the contribution of the present project.")
    add_sec_head("2.1 Conventional ultrasonic thickness measurement")
    add_p("A conventional pulse-echo thickness gauge introduces ultrasound into a material and detects echoes from material boundaries. Thickness is inferred from the travel time of the ultrasonic wave and the sound velocity in the material. Probe selection, coupling, surface condition, checking readings against known values and signal quality affect the measurement.")
    page_break()

    # =========================================================================
    # PAGE 14: 2.2 BASE PAPER — PAN ET AL. (2024)
    # =========================================================================
    add_sec_head("2.2 BASE PAPER — PAN ET AL. (2024)", space_before=16, space_after=6)
    p_cit = doc.add_paragraph()
    p_cit.paragraph_format.line_spacing = 1.35
    p_cit.paragraph_format.space_after = Pt(8)
    r = p_cit.add_run("L. Pan, K. Shi, L. Han, D. Qu, Y. Zhang and W. Chen, “Research on Wireless Passive Ultrasonic Thickness Measurement Technology Based on Pulse Compression Method,” Sensors, vol. 24, no. 24, article 8023, 2024. DOI: 10.3390/s24248023.")
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)

    add_sec_head("Research focus")
    add_p("The paper investigates wireless passive ultrasonic thickness measurement and the use of pulse compression to improve the quality of received ultrasonic echoes. It considers wireless energy-transfer concepts and ultrasonic thickness measurement, supported by modelling and experiments.")
    add_sec_head("Relevance to this project")
    add_p("It is the principal reference because it directly addresses wireless ultrasonic thickness measurement. It provides a research basis for discussing the combination of a measurement probe, wireless operation and thickness estimation.")
    add_sec_head("Difference from the present prototype")
    add_p("The paper uses dedicated ultrasonic thickness-measurement hardware and signal processing. The present prototype uses an HC-SR04 to measure air distance and estimates thickness using the reference height. The project should cite the paper as a base reference, not claim to reproduce its pulse-compression method.")
    page_break()

    # =========================================================================
    # PAGE 15: 2.3 GONG ET AL. (2023)
    # =========================================================================
    add_sec_head("2.3 GONG ET AL. (2023)", space_before=16, space_after=6)
    p_cit2 = doc.add_paragraph()
    p_cit2.paragraph_format.line_spacing = 1.35
    p_cit2.paragraph_format.space_after = Pt(8)
    r = p_cit2.add_run("W. Gong, X. Wang, Z. Yang, Z. Zhai, W. Feng and D. Liu, “Ultrasonic Thickness Measurement Method and System Implementation Based on Sampling Reconstruction and Phase Feature Extraction,” Sensors, vol. 23, no. 22, article 9072, 2023. DOI: 10.3390/s23229072.")
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)

    add_sec_head("Main theme")
    add_p("This study examines sampled ultrasonic signals and phase-feature extraction for thickness measurement. It illustrates the role of signal reconstruction and processing when the measurement is derived from ultrasonic echoes.")
    add_sec_head("Relation to the prototype")
    add_p("The current prototype does not implement sampling reconstruction or phase-feature extraction. It uses a commercial distance module and a subtraction formula. The paper is therefore a background reference for future work involving a dedicated ultrasonic probe and echo processing.")
    page_break()

    # =========================================================================
    # PAGE 16: 2.4 HE ET AL. (2025)
    # =========================================================================
    add_sec_head("2.4 HE ET AL. (2025)", space_before=16, space_after=6)
    p_cit3 = doc.add_paragraph()
    p_cit3.paragraph_format.line_spacing = 1.35
    p_cit3.paragraph_format.space_after = Pt(8)
    r = p_cit3.add_run("Y. He, Y. Chen, J. Wei, Z. Li, X. Guo, R. Xie, R. Li, J. Liu, Z. Tan and K. Zhang, “Design and Application of Wireless Wall Thickness Monitoring System for Ground Testing Process,” Processes, vol. 13, no. 1, article 63, 2025. DOI: 10.3390/pr13010063.")
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)

    add_sec_head("Research focus")
    add_p("The study describes a wireless wall-thickness monitoring system for ground-testing applications. It is relevant to remote reading collection, communication between measurement points and a monitoring interface.")
    add_sec_head("Relevance and limitation")
    add_p("The paper helps frame system-level design choices for wireless monitoring. Its sensor technology and industrial test arrangement differ from the HC-SR04 setup in the present prototype. It should be used as a supporting reference rather than as proof that the current hardware performs industrial wall-thickness measurement.")
    page_break()

    # =========================================================================
    # PAGE 17: 2.5 COMPARISON OF SELECTED STUDIES
    # =========================================================================
    add_sec_head("2.5 COMPARISON OF SELECTED STUDIES", space_before=16, space_after=8)
    lit_tbl = [
        ("Pan et al. (2024)", "Wireless passive ultrasonic thickness measurement", "Closest match to project theme", "Uses dedicated ultrasonic measurement and pulse compression"),
        ("Gong et al. (2023)", "Sampling reconstruction and phase features", "Signal-processing background", "Current firmware does not implement this method"),
        ("He et al. (2025)", "Wireless wall-thickness monitoring", "Wireless system layout", "Uses a different setup"),
        ("Current prototype", "Distance-based estimate with wireless display", "Low-cost embedded integration", "HC-SR04 measures distance through air")
    ]
    tbl_comp = doc.add_table(rows=len(lit_tbl)+1, cols=4)
    tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_full_grid_borders(tbl_comp)
    tbl_comp.rows[0].cells[0].width = Inches(1.5)
    tbl_comp.rows[0].cells[1].width = Inches(1.8)
    tbl_comp.rows[0].cells[2].width = Inches(1.6)
    tbl_comp.rows[0].cells[3].width = Inches(1.8)

    for c_i, h_txt in enumerate(["Study", "Main focus", "Relevance", "Difference from prototype"]):
        cell = tbl_comp.rows[0].cells[c_i]
        cell.paragraphs[0].text = h_txt
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "E2E8F0")

    for idx, (st, mf, rel, diff) in enumerate(lit_tbl):
        r_c = tbl_comp.rows[idx+1]
        r_c.cells[0].paragraphs[0].text = st
        r_c.cells[0].paragraphs[0].runs[0].bold = True
        r_c.cells[1].paragraphs[0].text = mf
        r_c.cells[2].paragraphs[0].text = rel
        r_c.cells[3].paragraphs[0].text = diff
        for cell in r_c.cells:
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
    page_break()

    # =========================================================================
    # PAGE 18: 2.6 RESEARCH GAP AND PROJECT POSITION
    # =========================================================================
    add_sec_head("2.6 RESEARCH GAP AND PROJECT POSITION", space_before=16, space_after=8)
    add_sec_head("Observations from literature")
    add_bullet("Thickness measurement depends on the sensor principle and checking readings against known values method.")
    add_bullet("Wireless monitoring can make readings available away from the sensing point.")
    add_bullet("Signal processing and probe design are central to conventional ultrasonic thickness gauging.")
    add_bullet("Different operating environments may require specialised transducers and mounting arrangements.")
    add_sec_head("Project position")
    add_p("The current project focuses on integration of distance sensing, embedded calculation, wireless transfer and local/remote display. It does not establish industrial ultrasonic wall-thickness capability. A future design could incorporate a dedicated ultrasonic probe, pulser/receiver circuitry, time-of-flight processing and checking readings against known values standards.")
    page_break()

    # =========================================================================
    # PAGE 19: 2.7 LITERATURE REVIEW SUMMARY
    # =========================================================================
    add_sec_head("2.7 LITERATURE REVIEW SUMMARY", space_before=16, space_after=8)
    add_p("The reviewed papers demonstrate that ultrasonic thickness measurement is an established engineering field with approaches for echo processing, wireless operation and specialised inspection applications. Pan et al. (2024) is selected as the base paper because it directly addresses wireless passive ultrasonic thickness measurement.")
    add_p("The present prototype uses a simpler and different method: it measures air distance with an HC-SR04, calculates a thickness estimate and transmits the value wirelessly. The report therefore separates the literature-based research context from the functions actually implemented.")
    add_sec_head("Design implications")
    add_bullet("State the working method accurately.")
    add_bullet("Validate readings against known-thickness samples.")
    add_bullet("Report the wireless function separately from measurement accuracy.")
    add_bullet("Match future hardware changes to the project title and base paper.")
    page_break()

    # =========================================================================
    # PAGE 20: CHAPTER 3 — SYSTEM REQUIREMENTS AND DESIGN
    # =========================================================================
    add_chapter_head("3", "System Requirements and Design")
    add_sec_head("3.1 Functional requirements")
    add_bullet("Acquire the distance to the top surface of an object.")
    add_bullet("Calculate estimated thickness using the configured reference height.")
    add_bullet("Transmit the calculated value to a separate receiver.")
    add_bullet("Display the received value on an LCD.")
    add_bullet("Make measurement data available through Blynk when configured.")
    add_sec_head("3.2 Non-functional requirements")
    add_bullet("Stable sensor mounting and known reference height.")
    add_bullet("Readable display and clear measurement units.")
    add_bullet("Reliable operation of the wireless link under test conditions.")
    add_bullet("Safe, regulated power for the electronic modules.")
    add_bullet("Firmware that rejects invalid readings and reports communication status.")
    page_break()

    # =========================================================================
    # PAGE 21: 3.3 SYSTEM ARCHITECTURE
    # =========================================================================
    add_sec_head("3.3 SYSTEM ARCHITECTURE", space_before=16, space_after=8)
    p_arch_flow = doc.add_paragraph()
    p_arch_flow.paragraph_format.space_before = Pt(8)
    p_arch_flow.paragraph_format.space_after = Pt(8)
    r = p_arch_flow.add_run("OBJECT SURFACE  →  HC-SR04  →  SENDER NODEMCU  →  ESP-NOW  →  RECEIVER NODEMCU  →  LCD\nREMOTE VIEWING THROUGH BLYNK IOT")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    r.font.color.rgb = NAVY_BLUE

    add_sec_head("Sender unit")
    add_p("The sender unit reads the distance sensor, applies the reference-height calculation and prepares a data packet. The controller is a NodeMCU ESP8266, as shown in the prototype photograph.")
    add_sec_head("Wireless link")
    add_p("ESP-NOW is used for direct communication between compatible ESP8266 devices. Both devices must use matching peer and channel settings.")
    add_sec_head("Receiver and display")
    add_p("The receiver interprets the received data and updates the 16×2 I2C LCD. Blynk provides an additional remote display path where network connectivity and application configuration are available.")
    page_break()

    # =========================================================================
    # PAGE 22: 3.4 THICKNESS CALCULATION
    # =========================================================================
    add_sec_head("3.4 THICKNESS CALCULATION", space_before=16, space_after=8)
    add_sec_head("Geometric relationship")
    p_formula = doc.add_paragraph()
    p_formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_formula.paragraph_format.space_before = Pt(10)
    p_formula.paragraph_format.space_after = Pt(10)
    r = p_formula.add_run("t = H − d")
    r.bold = True
    r.font.size = Pt(16)
    r.font.name = "Times New Roman"
    r.font.color.rgb = NAVY_BLUE

    add_p("Where t is the estimated thickness, H is the sensor reference height above the reference plane, and d is the distance measured from the sensor to the upper surface of the object.")
    add_p("The project setup states a reference height of H = 14.1 cm. The calculation assumes the object rests on the intended reference plane, the lower surface is aligned with that plane and the sensor position does not change.")
    add_sec_head("Illustrative arithmetic example")
    add_p("If H = 14.1 cm and a hypothetical measured distance is d = 10.0 cm, then t = 14.1 − 10.0 = 4.1 cm. This is only an example of the formula and is not an experimental result.")
    page_break()

    # =========================================================================
    # PAGE 23: 3.5 MEASUREMENT ERROR CONSIDERATIONS
    # =========================================================================
    add_sec_head("3.5 MEASUREMENT ERROR CONSIDERATIONS", space_before=16, space_after=8)
    add_p("The estimated thickness depends on both the measured distance and the reference height. If the sensor is moved after the reference height is recorded, the calculation can become systematically incorrect. Similarly, if an object is not placed on the reference plane, its apparent thickness may not match its actual thickness.")
    add_sec_head("Potential error sources")
    add_bullet("Incorrect reference-height measurement.")
    add_bullet("Object not aligned with the reference plane.")
    add_bullet("Sensor tilted relative to the target surface.")
    add_bullet("Target too small, angled or irregular for stable echo detection.")
    add_bullet("Distance readings fluctuate between repeated trials.")
    add_bullet("Incorrect unit conversion or stale wireless data.")
    add_sec_head("Good practice")
    add_p("Keep the mounting rigid, verify the reference height before testing, use consistent target placement and retain raw readings. Quantitative accuracy should be calculated from repeated comparisons with known reference values.")
    page_break()

    # =========================================================================
    # PAGE 24: 3.6 HARDWARE COMPONENTS
    # =========================================================================
    add_sec_head("3.6 HARDWARE COMPONENTS", space_before=16, space_after=8)
    hw_components = [
        ("HC-SR04", "Measures distance through air", "Target alignment and valid echo are important"),
        ("NodeMCU ESP8266 sender", "Acquisition and calculation", "Use firmware/pin mapping matching the board"),
        ("NodeMCU ESP8266 receiver", "Receives data and controls display", "Check wireless peer settings"),
        ("16×2 I2C LCD", "Local display", "Verify I2C address and wiring"),
        ("ESP-NOW", "Direct wireless data link", "Test communication separately"),
        ("Blynk IoT", "Remote display", "Requires correct configuration/connectivity"),
        ("L7805 regulator", "Nominal 5 V regulation", "Check current and heat dissipation")
    ]
    tbl_hw = doc.add_table(rows=len(hw_components)+1, cols=3)
    tbl_hw.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_full_grid_borders(tbl_hw)
    tbl_hw.rows[0].cells[0].width = Inches(2.0)
    tbl_hw.rows[0].cells[1].width = Inches(2.2)
    tbl_hw.rows[0].cells[2].width = Inches(2.5)

    for c_i, h_txt in enumerate(["Component", "Function", "Implementation note"]):
        cell = tbl_hw.rows[0].cells[c_i]
        cell.paragraphs[0].text = h_txt
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "E2E8F0")

    for idx, (comp, fn, note) in enumerate(hw_components):
        r_c = tbl_hw.rows[idx+1]
        r_c.cells[0].paragraphs[0].text = comp
        r_c.cells[0].paragraphs[0].runs[0].bold = True
        r_c.cells[1].paragraphs[0].text = fn
        r_c.cells[2].paragraphs[0].text = note
        for cell in r_c.cells:
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
    page_break()

    # =========================================================================
    # PAGE 25: 3.7 SOFTWARE AND COMMUNICATION REQUIREMENTS
    # =========================================================================
    add_sec_head("3.7 SOFTWARE AND COMMUNICATION REQUIREMENTS", space_before=16, space_after=8)
    sw_reqs = [
        ("Arduino IDE", "Compile and upload firmware", "Successful upload and serial output"),
        ("Embedded C/C++", "Sensor, calculation and display logic", "Inspect readings and error handling"),
        ("ESP-NOW", "Sender-to-receiver transfer", "Confirm packets are received"),
        ("LCD library", "Display control", "Verify address and text format"),
        ("Blynk", "Remote dashboard", "Confirm valid dashboard update"),
        ("Serial Monitor", "Debugging", "Compare raw distance and calculated estimate")
    ]
    tbl_sw = doc.add_table(rows=len(sw_reqs)+1, cols=3)
    tbl_sw.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_full_grid_borders(tbl_sw)
    tbl_sw.rows[0].cells[0].width = Inches(1.8)
    tbl_sw.rows[0].cells[1].width = Inches(2.4)
    tbl_sw.rows[0].cells[2].width = Inches(2.5)

    for c_i, h_txt in enumerate(["Software / protocol", "Purpose", "Verification"]):
        cell = tbl_sw.rows[0].cells[c_i]
        cell.paragraphs[0].text = h_txt
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "E2E8F0")

    for idx, (sw, purp, ver) in enumerate(sw_reqs):
        r_c = tbl_sw.rows[idx+1]
        r_c.cells[0].paragraphs[0].text = sw
        r_c.cells[0].paragraphs[0].runs[0].bold = True
        r_c.cells[1].paragraphs[0].text = purp
        r_c.cells[2].paragraphs[0].text = ver
        for cell in r_c.cells:
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    p_sw_note = doc.add_paragraph()
    p_sw_note.paragraph_format.space_before = Pt(12)
    p_sw_note.paragraph_format.line_spacing = 1.35
    r = p_sw_note.add_run("Record the actual board version, firmware settings, library versions and pin assignments used in the final build. Do not include private IoT authentication tokens in the printed report.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    page_break()

    # =========================================================================
    # PAGE 26: CHAPTER 4 — HARDWARE DESCRIPTION
    # =========================================================================
    add_chapter_head("4", "Hardware Description")
    add_sec_head("4.1 HC-SR04 distance sensor")
    add_p("The HC-SR04 is a low-cost ultrasonic distance module designed to estimate the distance to a reflecting surface in air. It uses trigger and echo signals. The controller initiates a measurement and determines the echo pulse duration, from which the module's distance estimate is calculated.")
    add_sec_head("Role in this project")
    add_p("The sensor is used to measure the air gap between its mounting position and the top surface of an object. The measured distance is not a material-wall thickness measurement. Thickness is estimated only because the mounting height is treated as a known reference.")
    page_break()

    # =========================================================================
    # PAGE 27: 4.2 NODEMCU ESP8266 CONTROLLER
    # =========================================================================
    add_sec_head("4.2 NODEMCU ESP8266 CONTROLLER", space_before=16, space_after=8)
    add_p("The NodeMCU ESP8266 is a development board built around the ESP8266 Wi-Fi-enabled microcontroller. It provides digital input/output and supports embedded firmware for reading sensors, performing calculations and exchanging data wirelessly.")
    add_sec_head("Sender responsibilities")
    add_bullet("Read the sensor output.")
    add_bullet("Check whether the reading is valid.")
    add_bullet("Calculate the estimated thickness.")
    add_bullet("Prepare the value for wireless transmission.")
    add_bullet("Provide serial output for debugging.")
    add_sec_head("Receiver responsibilities")
    add_bullet("Receive the sender's packet.")
    add_bullet("Validate the received data.")
    add_bullet("Update the LCD display.")
    add_bullet("Support the remote monitoring workflow where implemented.")
    page_break()

    # =========================================================================
    # PAGE 28: 4.3 ESP-NOW WIRELESS COMMUNICATION
    # =========================================================================
    add_sec_head("4.3 ESP-NOW WIRELESS COMMUNICATION", space_before=16, space_after=8)
    add_p("ESP-NOW is a connectionless communication protocol supported by Espressif devices. It can exchange small data packets directly between compatible devices without requiring a conventional Wi-Fi router for the direct sender-receiver link. Correct peer configuration is essential.")
    add_sec_head("Integration checks")
    add_bullet("Confirm the sender and receiver are compatible ESP8266 boards.")
    add_bullet("Verify the wireless channel and peer configuration.")
    add_bullet("Check the received value against the sender's serial output.")
    add_bullet("Test communication at the intended operating separation.")
    add_bullet("Record lost or delayed updates rather than assuming perfect delivery.")
    page_break()

    # =========================================================================
    # PAGE 29: 4.4 LCD AND BLYNK MONITORING
    # =========================================================================
    add_sec_head("4.4 LCD AND BLYNK MONITORING", space_before=16, space_after=8)
    add_sec_head("LCD display")
    add_p("A 16×2 I2C LCD provides local display of the calculated value. The I2C interface reduces the number of signal wires required compared with a parallel LCD connection. The display should use clear labels and units so the operator can interpret the reading.")
    add_sec_head("Blynk monitoring")
    add_p("Blynk is used as an additional remote interface for viewing measurement data. It requires the correct device and dashboard configuration and suitable network connectivity. Remote display should be tested separately from the LCD to identify where a fault occurs.")
    add_sec_head("Recommended display states")
    add_bullet("Valid reading with units.")
    add_bullet("No valid echo or invalid measurement.")
    add_bullet("Wireless connection unavailable or data not recently updated.")
    page_break()

    # =========================================================================
    # PAGE 30: 4.5 POWER SUPPLY AND SAFETY
    # =========================================================================
    add_sec_head("4.5 POWER SUPPLY AND SAFETY", space_before=16, space_after=8)
    add_p("The project documentation describes a 12 V source and L7805 voltage regulation. A linear regulator provides a nominal 5 V output but dissipates heat according to the voltage drop and load current. Wiring and regulator ratings must be checked before operating the circuit.")
    add_sec_head("Checks before powering")
    add_bullet("Verify supply polarity and connections.")
    add_bullet("Measure the regulated output with a multimeter.")
    add_bullet("Check the voltage requirements of each board and peripheral.")
    add_bullet("Do not connect 12 V directly to a 5 V input.")
    add_bullet("Ensure the regulator does not overheat under load.")
    add_bullet("Secure exposed conductors and prevent accidental short circuits.")
    p_pwr_note = doc.add_paragraph()
    p_pwr_note.paragraph_format.space_before = Pt(12)
    p_pwr_note.paragraph_format.line_spacing = 1.35
    r = p_pwr_note.add_run("The final report should reflect the actual supply wiring used in the prototype. If the implemented circuit differs from the original design description, update this section before submission.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    page_break()

    # =========================================================================
    # PAGE 31: CHAPTER 5 — METHODOLOGY
    # =========================================================================
    add_chapter_head("5", "Methodology")
    meth_steps = [
        "1. Review related literature and identify the functions required for the prototype.",
        "2. Prepare the block diagram and divide the sender and receiver responsibilities.",
        "3. Mount the HC-SR04 at the selected reference height.",
        "4. Interface the sensor with the sender NodeMCU.",
        "5. Develop firmware to read distance and calculate the estimate.",
        "6. Configure ESP-NOW and transmit the measurement value.",
        "7. Program the receiver to display the received value on the LCD.",
        "8. Configure and test the Blynk monitoring interface.",
        "9. Perform functional tests and compare readings with known-thickness samples."
    ]
    for ms in meth_steps:
        add_bullet(ms, space_after=6)
    page_break()

    # =========================================================================
    # PAGE 32: 5.2 SENDER UNIT WORKING
    # =========================================================================
    add_sec_head("5.2 SENDER UNIT WORKING", space_before=16, space_after=8)
    add_p("At the start of a measurement cycle, the sender triggers the HC-SR04 module. The module emits an ultrasonic burst and detects its reflection from the target. The measured echo duration is used to estimate the air distance. The controller then subtracts the distance from the configured reference height.")
    add_sec_head("Calculation and validation")
    add_p("The firmware should ensure that the distance value is valid and within the intended measurement geometry. If the sensor does not receive a suitable echo, the program should indicate an invalid reading instead of generating a misleading thickness estimate.")
    add_sec_head("Data preparation")
    add_p("After calculation, the sender packages the value for transmission. The data format should be consistent between sender and receiver, including the numeric type and units. Debugging output helps compare the measured distance with the transmitted value.")
    page_break()

    # =========================================================================
    # PAGE 33: 5.3 RECEIVER UNIT WORKING
    # =========================================================================
    add_sec_head("5.3 RECEIVER UNIT WORKING", space_before=16, space_after=8)
    add_p("The receiver listens for ESP-NOW packets from the sender. When a packet is received, the firmware checks its content and updates the display. The LCD value should correspond to the most recently received valid measurement.")
    add_sec_head("Receiver checks")
    add_bullet("Confirm that the receiver starts correctly.")
    add_bullet("Verify that packets arrive from the intended sender.")
    add_bullet("Compare LCD values with sender-side data.")
    add_bullet("Check the display formatting and units.")
    add_bullet("Test what the system displays when packets stop arriving.")
    add_p("A future version can add a packet counter or timestamp so that old readings can be distinguished from current readings.")
    page_break()

    # =========================================================================
    # PAGE 34: 5.4 INTEGRATION AND TEST SEQUENCE
    # =========================================================================
    add_sec_head("5.4 INTEGRATION AND TEST SEQUENCE", space_before=16, space_after=8)
    test_stages = [
        ("1. Power", "Verify regulated supply", "Boards start without reset or overheating"),
        ("2. Sensing", "Measure a stable target", "Distance value is plausible and repeatable"),
        ("3. Calculation", "Apply H − d", "Estimate matches manual calculation"),
        ("4. Wireless", "Send a value", "Receiver obtains the same value"),
        ("5. LCD", "Update display", "Value and units are readable"),
        ("6. IoT", "Update dashboard", "Blynk shows the transmitted value"),
        ("7. Validation", "Compare to reference", "Error can be calculated from actual readings")
    ]
    tbl_tseq = doc.add_table(rows=len(test_stages)+1, cols=3)
    tbl_tseq.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_full_grid_borders(tbl_tseq)
    tbl_tseq.rows[0].cells[0].width = Inches(1.8)
    tbl_tseq.rows[0].cells[1].width = Inches(2.4)
    tbl_tseq.rows[0].cells[2].width = Inches(2.5)

    for c_i, h_txt in enumerate(["Stage", "Action", "Expected check"]):
        cell = tbl_tseq.rows[0].cells[c_i]
        cell.paragraphs[0].text = h_txt
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "E2E8F0")

    for idx, (stg, act, exp) in enumerate(test_stages):
        r_c = tbl_tseq.rows[idx+1]
        r_c.cells[0].paragraphs[0].text = stg
        r_c.cells[0].paragraphs[0].runs[0].bold = True
        r_c.cells[1].paragraphs[0].text = act
        r_c.cells[2].paragraphs[0].text = exp
        for cell in r_c.cells:
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    p_tseq_note = doc.add_paragraph()
    p_tseq_note.paragraph_format.space_before = Pt(12)
    p_tseq_note.paragraph_format.line_spacing = 1.35
    r = p_tseq_note.add_run("Complete each stage before moving to the next. This staged process makes troubleshooting more efficient and produces clearer evidence for the project review.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    page_break()

    # =========================================================================
    # PAGE 35: CHAPTER 6 — EXPERIMENTAL SETUP
    # =========================================================================
    add_chapter_head("6", "Experimental Setup")
    add_p("The sensor is mounted at the reference height of 14.1 cm above a stable reference plane. The object is placed below the sensor with its lower surface aligned to that plane. The sensing unit and receiver are powered, and the LCD is checked before data collection.")
    add_sec_head("Equipment required")
    add_bullet("Assembled sender and receiver units.")
    add_bullet("Known-thickness samples or a suitable reference measuring instrument.")
    add_bullet("Ruler or height gauge to verify the sensor height.")
    add_bullet("Multimeter to check the regulated supply.")
    add_bullet("Computer with the firmware development environment.")
    add_bullet("Mobile device with the configured Blynk dashboard.")
    page_break()

    # =========================================================================
    # PAGE 36: 6.2 EXPERIMENTAL PROCEDURE
    # =========================================================================
    add_sec_head("6.2 EXPERIMENTAL PROCEDURE", space_before=16, space_after=8)
    exp_steps = [
        "1. Inspect the circuit and verify the power supply.",
        "2. Measure and record the sensor reference height.",
        "3. Power the sender and receiver.",
        "4. Place a known-thickness sample on the reference plane.",
        "5. Record the distance and calculated estimate.",
        "6. Confirm the same value appears on the receiver LCD.",
        "7. Check Blynk when network connectivity is available.",
        "8. Repeat each sample measurement several times.",
        "9. Compare the estimate with the reference thickness.",
        "10. Record any invalid readings or communication interruptions."
    ]
    for es in exp_steps:
        add_bullet(es, space_after=5)
    page_break()

    # =========================================================================
    # PAGE 37: 6.3 EXPERIMENTAL DATA
    # =========================================================================
    add_sec_head("6.3 EXPERIMENTAL DATA", space_before=16, space_after=8)
    add_p("Enter only values measured during the actual test. The blanks below are intentionally retained because no verified measurement readings was supplied for this report.")
    
    tbl_data = doc.add_table(rows=6, cols=4)
    tbl_data.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_full_grid_borders(tbl_data)
    for c_i, h_txt in enumerate(["Sample", "Reference thickness (cm)", "Measured distance (cm)", "Estimated thickness (cm)"]):
        cell = tbl_data.rows[0].cells[c_i]
        cell.paragraphs[0].text = h_txt
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "E2E8F0")

    for s_i in range(1, 6):
        r_c = tbl_data.rows[s_i]
        r_c.cells[0].paragraphs[0].text = str(s_i)
        r_c.cells[0].paragraphs[0].runs[0].bold = True
        r_c.cells[1].paragraphs[0].text = "________"
        r_c.cells[2].paragraphs[0].text = "________"
        r_c.cells[3].paragraphs[0].text = "________"
        for cell in r_c.cells:
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    p_data_sub = doc.add_paragraph()
    p_data_sub.paragraph_format.space_before = Pt(12)
    p_data_sub.paragraph_format.line_spacing = 1.35
    r = p_data_sub.add_run("Estimated thickness = 14.1 cm − measured distance, provided the reference height remains 14.1 cm and the object is aligned with the reference plane.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    page_break()

    # =========================================================================
    # PAGE 38: 6.4 ERROR CALCULATION AND CALIBRATION
    # =========================================================================
    add_sec_head("6.4 ERROR CALCULATION AND CALIBRATION", space_before=16, space_after=8)
    p_err_f = doc.add_paragraph()
    p_err_f.paragraph_format.line_spacing = 1.35
    p_err_f.paragraph_format.space_after = Pt(8)
    r1 = p_err_f.add_run("Signed error = estimated thickness − reference thickness\n")
    r2 = p_err_f.add_run("Absolute error = |estimated thickness − reference thickness|\n")
    r3 = p_err_f.add_run("Percentage error = 100 × absolute error / reference thickness")
    for r in [r1, r2, r3]:
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(11.5)
        r.font.color.rgb = NAVY_BLUE

    add_p("Percentage error is applicable only when the reference thickness is non-zero. For repeat measurements, calculate the mean estimated value and a measure of spread. State the number of trials, reference instrument and test conditions.")
    add_sec_head("Calibration notes")
    add_bullet("Verify the reference height before each test series.")
    add_bullet("Use known-thickness samples with independently known thickness.")
    add_bullet("Repeat measurements using consistent sample placement.")
    add_bullet("Keep raw distance readings as well as calculated values.")
    add_bullet("Do not report an accuracy percentage until actual data have been analysed.")
    page_break()

    # =========================================================================
    # PAGE 39: CHAPTER 7 — RESULTS AND DISCUSSION
    # =========================================================================
    add_chapter_head("7", "Results and Discussion")
    add_p("The project documentation describes working connection of the main parts of distance acquisition, thickness estimation using the reference height, ESP-NOW transmission, local LCD display and Blynk monitoring. These are functions described in the project. The supplied material does not contain a verified numerical dataset from which measurement accuracy or consistency of repeated readings can be calculated.")
    
    res_table_data = [
        ("Distance sensor", "Distance acquisition", "Raw distance readings"),
        ("Calculation", "Reference height minus distance", "Calculation check"),
        ("ESP-NOW", "Wireless transfer", "Sender/receiver comparison"),
        ("LCD", "Local display", "Photograph or observation record"),
        ("Blynk", "Remote monitoring", "Dashboard screenshot or log")
    ]
    tbl_res = doc.add_table(rows=len(res_table_data)+1, cols=3)
    tbl_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_full_grid_borders(tbl_res)
    tbl_res.rows[0].cells[0].width = Inches(2.0)
    tbl_res.rows[0].cells[1].width = Inches(2.3)
    tbl_res.rows[0].cells[2].width = Inches(2.4)

    for c_i, h_txt in enumerate(["Subsystem", "Function considered", "Evidence to retain"]):
        cell = tbl_res.rows[0].cells[c_i]
        cell.paragraphs[0].text = h_txt
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "E2E8F0")

    for idx, (sub, fn, ev) in enumerate(res_table_data):
        r_c = tbl_res.rows[idx+1]
        r_c.cells[0].paragraphs[0].text = sub
        r_c.cells[0].paragraphs[0].runs[0].bold = True
        r_c.cells[1].paragraphs[0].text = fn
        r_c.cells[2].paragraphs[0].text = ev
        for cell in r_c.cells:
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
    page_break()

    # =========================================================================
    # PAGE 40: 7.2 DISCUSSION
    # =========================================================================
    add_sec_head("7.2 DISCUSSION", space_before=16, space_after=8)
    add_p("The system sends the calculated value wirelessly from the sender to the receiver. The estimate depends on the 14.1 cm reference height and correct object placement. If the sensor or object moves out of position, the reading may change.")
    add_sec_head("Interpreting the prototype")
    add_bullet("Wireless transmission demonstrates communication, not measurement accuracy.")
    add_bullet("A working display confirms that a value reached the receiver, not that the value is physically correct.")
    add_bullet("Reference comparison is needed to evaluate measurement error.")
    add_bullet("Repeated trials are needed to assess consistency of repeated readings.")
    add_bullet("The sensor's working method must be described accurately.")
    page_break()

    # =========================================================================
    # PAGE 41: 7.3 LIMITATIONS
    # =========================================================================
    add_sec_head("7.3 LIMITATIONS", space_before=16, space_after=8)
    add_bullet("The HC-SR04 measures air distance to an external surface.")
    add_bullet("The current method estimates thickness from a fixed geometric reference.")
    add_bullet("Object position, surface angle and sensor alignment can affect readings.")
    add_bullet("No verified accuracy or consistency of repeated readings values are available in the supplied report.")
    add_bullet("The current arrangement is not equivalent to an industrial ultrasonic wall-thickness gauge.")
    add_sec_head("Recommended next steps")
    add_p("Complete the known-thickness sample tests, calculate the errors and include a small table of actual findings. For a true material-thickness gauge, the next design stage would require a suitable ultrasonic probe, pulser/receiver electronics, time-of-flight processing and checking readings against known values standards.")
    page_break()

    # =========================================================================
    # PAGE 42: CHAPTER 8 — CONCLUSION AND FUTURE SCOPE
    # =========================================================================
    add_chapter_head("8", "Conclusion and Future Scope")
    add_sec_head("8.1 Conclusion")
    add_p("A wireless prototype was assembled using an HC-SR04 distance sensor, NodeMCU ESP8266 controllers, ESP-NOW communication, a 16×2 I2C LCD and Blynk monitoring. The system estimates object thickness by subtracting the measured air distance from a fixed reference height of 14.1 cm.")
    add_p("The prototype demonstrates integration of sensing, embedded calculation, wireless data transfer and local/remote display. Since no verified numerical test dataset was supplied, this report does not claim a specific accuracy, error percentage or consistency of repeated readings.")
    add_sec_head("8.2 Future scope")
    add_bullet("Validate the estimate against known known-thickness samples.")
    add_bullet("Add timestamped data logging and communication status.")
    add_bullet("Improve invalid-reading detection and consistency of repeated readings.")
    add_bullet("For material wall-thickness measurement, integrate a dedicated ultrasonic probe, pulser/receiver and time-of-flight processing.")
    page_break()

    # =========================================================================
    # PAGE 43: REFERENCES AND ACTUAL PROTOTYPE
    # =========================================================================
    add_page_title("REFERENCES AND ACTUAL PROTOTYPE", space_before=16, space_after=10)
    refs = [
        "[1] L. Pan, K. Shi, L. Han, D. Qu, Y. Zhang and W. Chen, “Research on Wireless Passive Ultrasonic Thickness Measurement Technology Based on Pulse Compression Method,” Sensors, vol. 24, no. 24, article 8023, 2024. DOI: 10.3390/s24248023.",
        "[2] W. Gong, X. Wang, Z. Yang, Z. Zhai, W. Feng and D. Liu, “Ultrasonic Thickness Measurement Method and System Implementation Based on Sampling Reconstruction and Phase Feature Extraction,” Sensors, vol. 23, no. 22, article 9072, 2023. DOI: 10.3390/s23229072.",
        "[3] Y. He, Y. Chen, J. Wei, Z. Li, X. Guo, R. Xie, R. Li, J. Liu, Z. Tan and K. Zhang, “Design and Application of Wireless Wall Thickness Monitoring System for Ground Testing Process,” Processes, vol. 13, no. 1, article 63, 2025. DOI: 10.3390/pr13010063."
    ]
    for rf in refs:
        p_rf = doc.add_paragraph()
        p_rf.paragraph_format.line_spacing = 1.25
        p_rf.paragraph_format.space_after = Pt(4)
        r = p_rf.add_run(rf)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)

    add_sec_head("Actual prototype photographs", space_before=8, space_after=4)
    fig1_path = os.path.abspath("prototype_fig_43_0.png")
    if os.path.exists(fig1_path):
        p_fig1 = doc.add_paragraph()
        p_fig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig1.paragraph_format.space_before = Pt(4)
        p_fig1.paragraph_format.space_after = Pt(2)
        p_fig1.add_run().add_picture(fig1_path, width=Inches(3.8))

    p_cap1 = doc.add_paragraph()
    p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap1.paragraph_format.space_after = Pt(4)
    r = p_cap1.add_run("Figure 1. Electronics and LCD assembly with NodeMCU ESP8266.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.italic = True
    page_break()

    # =========================================================================
    # PAGE 44: PROTOTYPE PHOTO 2
    # =========================================================================
    fig2_path = os.path.abspath("prototype_fig_44_0.png")
    if os.path.exists(fig2_path):
        p_fig2 = doc.add_paragraph()
        p_fig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig2.paragraph_format.space_before = Pt(14)
        p_fig2.paragraph_format.space_after = Pt(4)
        p_fig2.add_run().add_picture(fig2_path, width=Inches(4.2))

    p_cap2 = doc.add_paragraph()
    p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap2.paragraph_format.space_after = Pt(12)
    r = p_cap2.add_run("Figure 2. Overall prototype arrangement showing the separated sensing and display/controller units.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.italic = True
    page_break()

    # =========================================================================
    # PAGE 45: APPENDIX — FINAL SUBMISSION CHECKLIST
    # =========================================================================
    add_page_title("APPENDIX — FINAL SUBMISSION CHECKLIST", space_before=16, space_after=12)
    add_p("Use this checklist before printing and submitting the report:")
    chk_items = [
        "Confirm that the approved project title is used consistently throughout the report.",
        "Verify all student names and register numbers on the cover and certificate.",
        "Obtain the required signature/approval from Dr. S. Kailainathan, Project Guide and Head of the Department.",
        "Check that the hardware description matches the assembled prototype.",
        "Insert actual measurement readings and reference values after testing.",
        "Calculate errors only from verified measurements; do not use the illustrative calculation as a result.",
        "Check that all prototype photographs are clear and correctly captioned.",
        "Verify reference details and DOI links before final submission.",
        "Update the table of contents and page numbers after the final formatting.",
        "Print one copy for review before printing the full set."
    ]
    for ci in chk_items:
        add_bullet(ci, space_after=6)

    # Save to both paths
    os.makedirs(os.path.dirname(DOCX_OUTPUT_PATH), exist_ok=True)
    doc.save(DOCX_OUTPUT_PATH)
    print(f"Saved DOCX: {DOCX_OUTPUT_PATH}")

    try:
        os.makedirs(os.path.dirname(DOCX_COPY_PATH), exist_ok=True)
        doc.save(DOCX_COPY_PATH)
        print(f"Saved copy to: {DOCX_COPY_PATH}")
    except Exception as e:
        print(f"Could not copy to {DOCX_COPY_PATH}: {e}")

if __name__ == '__main__':
    build_report()
